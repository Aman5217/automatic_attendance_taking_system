"""
reports.py - Generate Excel and CSV attendance reports
"""

import os
import pandas as pd
from datetime import datetime, date
from database import get_report_data, get_all_students, get_all_subjects

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def generate_session_report(session_id, fmt="xlsx"):
    """Detailed report for a single attendance session."""
    from database import get_attendance_by_session, get_session
    sess    = get_session(session_id)
    records = get_attendance_by_session(session_id)

    from database import get_subject
    subj = get_subject(sess["subject_id"])

    rows = []
    for r in records:
        rows.append({
            "Roll No":  r["roll_number"],
            "Name":     r["name"],
            "Department": r["department"],
            "Semester": r["semester"],
            "Section":  r["section"],
            "Status":   r["status"],
        })

    df = pd.DataFrame(rows)
    present = sum(1 for r in rows if r["Status"] == "Present")
    total   = len(rows)
    filename = f"Session_{session_id}_{sess['date']}_{subj['code']}.{fmt}"
    path = os.path.join(REPORTS_DIR, filename)

    if fmt == "csv":
        df.to_csv(path, index=False)
    else:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Attendance")
            _style_session_sheet(writer.sheets["Attendance"], rows, sess, subj, present, total)

    return path


def _style_session_sheet(ws, rows, sess, subj, present, total):
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    # Title rows
    ws.insert_rows(1, amount=5)
    ws["A1"] = "UNIVERSITY ATTENDANCE REPORT"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="2E4057")
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = f"Subject: {subj['name']} ({subj['code']})"
    ws["A3"] = f"Date: {sess['date']}  |  Time: {sess['start_time']}"
    ws["A4"] = f"Department: {sess['department']}  |  Semester: {sess['semester']}  |  Section: {sess['section']}"
    ws["A5"] = f"Present: {present}  |  Absent: {total-present}  |  Total: {total}  |  Attendance: {present/total*100:.1f}%" if total else ""

    for row_num in [2, 3, 4, 5]:
        ws[f"A{row_num}"].font = Font(size=10, color="2E4057")
        ws.merge_cells(f"A{row_num}:F{row_num}")

    # Header row
    header_fill = PatternFill("solid", fgColor="048A81")
    for cell in ws[6]:
        cell.fill   = header_fill
        cell.font   = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    green_fill = PatternFill("solid", fgColor="D4EDDA")
    red_fill   = PatternFill("solid", fgColor="F8D7DA")
    for i, row_data in enumerate(rows, start=7):
        fill = green_fill if row_data["Status"] == "Present" else red_fill
        for cell in ws[i]:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    # Column widths
    widths = [12, 28, 18, 10, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def generate_monthly_report(department, semester, section, subject_id=None,
                             start_date=None, end_date=None, fmt="xlsx"):
    """
    Full attendance matrix: rows=students, columns=sessions
    Shows P/A for each session with percentage.
    """
    data, sessions = get_report_data(department, semester, section, subject_id, start_date, end_date)

    if not sessions:
        return None

    from database import get_subject
    rows = []
    for sid, info in data.items():
        s = info["student"]
        rec = info["records"]
        present = sum(1 for v in rec.values() if v == "P")
        total   = len(sessions)
        pct     = round(present / total * 100, 1) if total else 0
        row = {
            "Roll No":  s["roll_number"],
            "Name":     s["name"],
        }
        for sess in sessions:
            col = f"{sess['date']}"
            row[col] = rec.get(sess["id"], "A")
        row["Present"] = present
        row["Total"]   = total
        row["Attendance %"] = f"{pct}%"
        row["Status"]  = "OK" if pct >= 75 else "SHORTAGE"
        rows.append(row)

    df = pd.DataFrame(rows)
    label = f"{start_date}_to_{end_date}" if start_date else date.today().isoformat()
    subj_tag = str(subject_id) if subject_id else "All"
    filename = f"Report_{department}_Sem{semester}_{section}_{label}_{subj_tag}.{fmt}"
    path = os.path.join(REPORTS_DIR, filename)

    if fmt == "csv":
        df.to_csv(path, index=False)
        return path

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Attendance Matrix")
        ws = writer.sheets["Attendance Matrix"]
        _style_matrix_sheet(ws, rows, df.columns.tolist(), sessions)

    return path


def _style_matrix_sheet(ws, rows, columns, sessions):
    from openpyxl.styles import PatternFill, Font, Alignment

    header_fill = PatternFill("solid", fgColor="2E4057")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(horizontal="center")

    green  = PatternFill("solid", fgColor="D4EDDA")
    red    = PatternFill("solid", fgColor="F8D7DA")
    yellow = PatternFill("solid", fgColor="FFF3CD")
    blue   = PatternFill("solid", fgColor="CCE5FF")

    for i, row_data in enumerate(rows, start=2):
        status = row_data.get("Status", "OK")
        row_fill = green if status == "OK" else yellow
        for j, cell in enumerate(ws[i]):
            val = cell.value
            if val == "P":
                cell.fill = green
                cell.font = Font(bold=True, color="155724")
            elif val == "A":
                cell.fill = red
                cell.font = Font(color="721C24")
            elif val == "SHORTAGE":
                cell.fill = yellow
                cell.font = Font(bold=True, color="856404")
            else:
                cell.fill = row_fill
            cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max_len, 20)


def generate_student_report(student_id, fmt="xlsx"):
    """Individual student attendance report across all subjects."""
    from database import get_student, get_student_sessions
    student  = get_student(student_id)
    sessions = get_student_sessions(student_id)

    rows = []
    for s in sessions:
        rows.append({
            "Date":    s["date"],
            "Subject": s["subject_name"],
            "Code":    s["code"],
            "Status":  s["status"],
            "Time":    s["start_time"],
        })

    df = pd.DataFrame(rows)
    filename = f"Student_{student['roll_number']}_{date.today()}.{fmt}"
    path = os.path.join(REPORTS_DIR, filename)

    if fmt == "csv":
        df.to_csv(path, index=False)
    else:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="My Attendance")
            ws = writer.sheets["My Attendance"]
            from openpyxl.styles import PatternFill, Font, Alignment
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="2E4057")
                cell.font = Font(bold=True, color="FFFFFF")
            green = PatternFill("solid", fgColor="D4EDDA")
            red   = PatternFill("solid", fgColor="F8D7DA")
            for row in ws.iter_rows(min_row=2):
                fill = green if row[3].value == "Present" else red
                for cell in row:
                    cell.fill = fill
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 18

    return path
